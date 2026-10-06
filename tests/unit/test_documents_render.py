"""DOCX and PDF renderers: the realism markers survive, and the bytes are reproducible."""

from __future__ import annotations

import io
import re
import zipfile
from typing import Any

import docx
import pytest
from pypdf import PdfReader

from csfd.documents.figures import raster_bytes, svg_to_png
from csfd.documents.ir import DocumentIR, Section
from csfd.documents.render_docx import render_docx
from csfd.documents.render_pdf import render_pdf
from tests.fixtures.sample_documents import SVG, figure_asset, service_manual

DOC = service_manual("Kalvora Dental")
ASSETS = {figure_asset().asset_id: figure_asset()}


@pytest.fixture(scope="module")
def docx_bytes() -> bytes:
    return render_docx(DOC, ASSETS)


@pytest.fixture(scope="module")
def pdf_render() -> tuple[bytes, dict[str, tuple[int, int]]]:
    return render_pdf(DOC, ASSETS)


def _xml(data: bytes, member: str) -> str:
    return zipfile.ZipFile(io.BytesIO(data)).read(member).decode("utf-8")


def test_docx_uses_real_heading_styles_with_typed_numbers(docx_bytes: bytes) -> None:
    d = docx.Document(io.BytesIO(docx_bytes))
    headings = [
        (p.style.name, p.text)
        for p in d.paragraphs
        if p.style is not None and p.style.name.startswith(("Heading", "Title"))
    ]
    assert ("Title", "CF-600 Ceramic Firing Furnace") in headings
    assert ("Heading 1", "6 Calibration") in headings
    assert ("Heading 2", "6.2 Silver test") in headings


def test_docx_warnings_are_styled_paragraphs_not_tables(docx_bytes: bytes) -> None:
    d = docx.Document(io.BytesIO(docx_bytes))
    by_style = {p.style.name: p.text for p in d.paragraphs if p.style is not None}
    assert by_style["Warning"].startswith("WARNING")
    assert by_style["Caution"].startswith("CAUTION")
    # Tables: document control, revision history, spare parts. No one-cell warning tables.
    assert len(d.tables) == 3


def test_docx_procedure_steps_are_numbered_with_results(docx_bytes: bytes) -> None:
    d = docx.Document(io.BytesIO(docx_bytes))
    steps = [
        p.text for p in d.paragraphs if p.style is not None and p.style.name == "List Paragraph"
    ]
    assert steps == [
        "1. Place the sample in the centre.",
        "2. Start Service > Calibration > Silver test. Result: The display shows the melt "
        "temperature.",
    ]


def test_docx_has_captions_bookmarks_alt_text_and_page_fields(docx_bytes: bytes) -> None:
    d = docx.Document(io.BytesIO(docx_bytes))
    captions = [p.text for p in d.paragraphs if p.style is not None and p.style.name == "Caption"]
    assert "Figure 3-1. CF-600 front view with item numbers" in captions
    assert "Table 9-1. Spare parts" in captions
    body = _xml(docx_bytes, "word/document.xml")
    bookmarks = re.findall(r'w:bookmarkStart w:id="\d+" w:name="([^"]+)"', body)
    assert bookmarks == ["sec_1", "sec_3", "sec_6", "sec_6__2", "sec_9"]
    assert 'descr="Line drawing of the CF-600 front' in body
    assert "<w:tblHeader" in body
    footer = _xml(docx_bytes, "word/footer1.xml")
    assert re.findall(r"<w:instrText[^>]*>([^<]+)<", footer) == [" PAGE ", " NUMPAGES "]
    assert d.core_properties.identifier == DOC.doc_id


def test_docx_bytes_are_reproducible(docx_bytes: bytes) -> None:
    assert render_docx(DOC, ASSETS) == docx_bytes


def test_pdf_has_outline_tags_metadata_and_page_x_of_y(
    pdf_render: tuple[bytes, dict[str, tuple[int, int]]],
) -> None:
    pdf, _ = pdf_render
    reader = PdfReader(io.BytesIO(pdf))
    root: Any = reader.trailer["/Root"]
    assert "/StructTreeRoot" in root  # tagged PDF
    assert reader.metadata is not None
    assert reader.metadata.title == "CF-600 Ceramic Firing Furnace - Service Manual (extract)"
    assert str(reader.metadata["/CreationDate"]).startswith("D:20250915")

    def titles(items: list[Any]) -> list[str]:
        out: list[str] = []
        for item in items:
            out += titles(item) if isinstance(item, list) else [item.title]
        return out

    assert titles(reader.outline) == [
        "1 Safety",
        "3 Main components",
        "6 Calibration",
        "6.2 Silver test",
        "9 Spare parts",
    ]
    last = reader.pages[-1].extract_text()
    assert f"Page {len(reader.pages)} of {len(reader.pages)}" in last
    assert "KD-60-2204" in last


def test_pdf_page_ranges_cover_every_section(
    pdf_render: tuple[bytes, dict[str, tuple[int, int]]],
) -> None:
    pdf, pages = pdf_render
    count = len(PdfReader(io.BytesIO(pdf)).pages)
    assert set(pages) == {s.id for s, _, _ in DOC.walk()}
    assert all(1 <= first <= last <= count for first, last in pages.values())


def test_pdf_bytes_are_reproducible(
    pdf_render: tuple[bytes, dict[str, tuple[int, int]]],
) -> None:
    assert render_pdf(DOC, ASSETS)[0] == pdf_render[0]


def test_svg_figures_become_png_for_docx() -> None:
    png = svg_to_png(SVG)
    assert png.startswith(b"\x89PNG")
    assert svg_to_png(SVG) == png
    assert raster_bytes(figure_asset()) == (png, "png")


def test_dotted_and_dashed_section_ids_get_distinct_bookmarks() -> None:
    sections = [
        Section(id="sec-6.2", number="6.2", title="A"),
        Section(id="sec-6-2", number="6.3", title="B"),
    ]
    doc = DocumentIR.model_validate({**DOC.model_dump(), "sections": sections})
    body = zipfile.ZipFile(io.BytesIO(render_docx(doc, {}))).read("word/document.xml").decode()
    bookmarks = re.findall(r'w:bookmarkStart w:id="\d+" w:name="([^"]+)"', body)
    assert bookmarks == ["sec_6__2", "sec_6_2"]
