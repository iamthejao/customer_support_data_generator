"""Render a document IR to DOCX with python-docx.

What makes the file look like a real manual, and read well in RAG ingesters
(Docling, Unstructured map Word styles to structure):

* real ``Title`` / ``Heading 1-3`` styles, with the section number typed into
  the heading text (ingesters do not all recompute Word auto-numbering);
* a bookmark per section named after its IR id (``sec-6.2`` -> ``sec_6__2``);
* safety messages as their own paragraph styles (``Danger``, ``Warning``, ...)
  with a coloured left border, not as one-cell tables;
* captioned tables with a repeated header row, captioned figures with alt text;
* a running header (company, document number, revision) and a footer with
  "Page X of Y" fields, which Word fills in when it lays the file out.

python-docx has no API for fields, bookmarks, borders or alt text, so those few
elements are written as OOXML. Output is byte-reproducible: core properties are
pinned to the issue date and the zip package is rewritten with fixed member
timestamps.
"""

from __future__ import annotations

import io
import zipfile
from collections.abc import Mapping
from datetime import UTC, datetime, time

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.styles.style import ParagraphStyle
from docx.text.paragraph import Paragraph as DocxParagraph

from csfd.documents.figures import raster_bytes
from csfd.documents.ir import (
    Admonition,
    Asset,
    DocumentIR,
    Figure,
    ItemList,
    Paragraph,
    Procedure,
    Section,
    Table,
    bookmark_name,
)

# Signal word and colour per safety level (ANSI Z535 colours, softened for print).
SIGNALS: dict[str, tuple[str, str]] = {
    "danger": ("DANGER", "C00000"),
    "warning": ("WARNING", "E36C09"),
    "caution": ("CAUTION", "BF8F00"),
    "notice": ("NOTICE", "1F4E79"),
    "note": ("NOTE", "595959"),
}
_FIGURE_WIDTH = Cm(14)
_TABLE_STYLE = "Light Grid Accent 1"
_STEP_STYLE = "List Paragraph"
_BULLET_STYLE = "List Bullet"


def render_docx(doc: DocumentIR, assets: Mapping[str, Asset]) -> bytes:
    """The document as DOCX bytes."""
    d = Document()
    _base_styles(d)
    _page_setup(d, doc)
    _front_matter(d, doc)
    bookmark_id = 0
    for section, depth, _ in doc.walk():
        bookmark_id += 1
        _section(d, section, depth, bookmark_id, assets)
    props = d.core_properties
    props.title = doc.title
    props.subject = doc.subtitle or ""
    props.identifier = doc.doc_id
    props.keywords = ";".join([doc.doc_type, doc.doc_number, *doc.applies_to.models])
    props.author = props.last_modified_by = doc.company
    props.language = doc.language
    props.revision = 1
    pinned = datetime.combine(doc.issue_date, time(0, 0), tzinfo=UTC)
    props.created = props.modified = props.last_printed = pinned
    buf = io.BytesIO()
    d.save(buf)
    return _normalise_zip(buf.getvalue(), pinned)


def _base_styles(d: DocxDocument) -> None:
    normal = d.styles["Normal"]
    assert isinstance(normal, ParagraphStyle)
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    for word, colour in SIGNALS.values():
        style = d.styles.add_style(word.title(), WD_STYLE_TYPE.PARAGRAPH)
        assert isinstance(style, ParagraphStyle)
        style.base_style = normal
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.left_indent = Cm(0.3)
        p_pr = style.element.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        for key, value in (("w:val", "single"), ("w:sz", "24"), ("w:space", "8")):
            left.set(qn(key), value)
        left.set(qn("w:color"), colour)
        borders.append(left)
        p_pr.append(borders)
        shading = OxmlElement("w:shd")
        for key, value in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", "F2F2F2")):
            shading.set(qn(key), value)
        p_pr.append(shading)
    title = d.styles.add_style("Procedure Title", WD_STYLE_TYPE.PARAGRAPH)
    assert isinstance(title, ParagraphStyle)
    title.base_style = normal
    title.font.bold = True
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.keep_with_next = True


def _page_setup(d: DocxDocument, doc: DocumentIR) -> None:
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)  # A4
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.2)
    header = sec.header.paragraphs[0]
    header.text = f"{doc.company}\t{doc.doc_number}\tRev. {doc.revision}"
    footer = sec.footer.paragraphs[0]
    footer.add_run(f"{doc.title}\tPage ")
    _field(footer, "PAGE")
    footer.add_run(" of ")
    _field(footer, "NUMPAGES")


def _front_matter(d: DocxDocument, doc: DocumentIR) -> None:
    d.add_paragraph(doc.title, style="Title")
    if doc.subtitle:
        d.add_paragraph(doc.subtitle, style="Subtitle")
    control = [
        ("Document number", doc.doc_number),
        ("Revision", doc.revision),
        ("Issue date", doc.issue_date.isoformat()),
        ("Status", doc.status),
        ("Applies to", ", ".join(doc.applies_to.models + doc.applies_to.serials) or "-"),
        ("Audience", doc.audience),
    ]
    table = d.add_table(rows=0, cols=2)
    table.style = _TABLE_STYLE
    for key, value in control:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = key, value
    if doc.revisions:
        d.add_heading("Revision history", level=1)
        _table(
            d,
            Table(
                id="revision-history",
                caption="Revision history",
                columns=["Rev.", "Date", "Description", "Author"],
                rows=[
                    [r.rev, r.issued.isoformat(), r.description, r.author] for r in doc.revisions
                ],
            ),
            caption=False,
        )


def _section(
    d: DocxDocument, section: Section, depth: int, bookmark_id: int, assets: Mapping[str, Asset]
) -> None:
    heading = d.add_heading(f"{section.number} {section.title}", level=min(depth, 3))
    _bookmark(heading, bookmark_name(section.id), bookmark_id)
    for block in section.blocks:
        if isinstance(block, Paragraph):
            d.add_paragraph(block.text)
        elif isinstance(block, Admonition):
            word, colour = SIGNALS[block.level]
            p = d.add_paragraph(style=word.title())
            run = p.add_run(f"{word}  ")
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(colour)
            p.add_run(block.text)
        elif isinstance(block, Procedure):
            d.add_paragraph(f"Procedure: {block.title}", style="Procedure Title")
            # Numbers are typed: Word list numbering would continue across procedures.
            for n, step in enumerate(block.steps, 1):
                p = d.add_paragraph(f"{n}. {step.text}", style=_STEP_STYLE)
                if step.expected:
                    p.add_run(f" Result: {step.expected}").italic = True
        elif isinstance(block, ItemList):
            for item in block.items:
                d.add_paragraph(item, style=_BULLET_STYLE)
        elif isinstance(block, Table):
            _table(d, block)
        elif isinstance(block, Figure):
            _figure(d, block, assets[block.asset_id])


def _table(d: DocxDocument, block: Table, *, caption: bool = True) -> None:
    if caption:
        d.add_paragraph(block.caption, style="Caption")
    table = d.add_table(rows=1, cols=len(block.columns))
    table.style = _TABLE_STYLE
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, name in zip(table.rows[0].cells, block.columns, strict=True):
        cell.text = name
    # Repeat the header row on every page the table spans.
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(header)
    for row in block.rows:
        for cell, value in zip(table.add_row().cells, row, strict=True):
            cell.text = value


def _figure(d: DocxDocument, block: Figure, asset: Asset) -> None:
    data, _ = raster_bytes(asset)
    d.add_picture(io.BytesIO(data), width=_FIGURE_WIDTH)
    picture = d.inline_shapes[-1]
    # wp:docPr/@descr is the alt text Word shows; @title names the figure.
    doc_pr = picture._inline.docPr
    doc_pr.set("descr", block.alt)
    doc_pr.set("title", block.id)
    d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_paragraph(block.caption, style="Caption")


def _field(paragraph: DocxParagraph, instruction: str) -> None:
    """A Word field (PAGE, NUMPAGES); "1" is the placeholder until Word updates it."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    for el in (begin, instr, separate):
        run._r.append(el)
    result = paragraph.add_run("1")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    result._r.append(end)


def _bookmark(paragraph: DocxParagraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _normalise_zip(data: bytes, pinned: datetime) -> bytes:
    """Rewrite the package with fixed member timestamps and order."""
    stamp = (pinned.year, pinned.month, pinned.day, 0, 0, 0)
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for name in sorted(src.namelist(), key=lambda n: (n != "[Content_Types].xml", n)):
            info = zipfile.ZipInfo(name, date_time=stamp)
            info.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(info, src.read(name))
    return out.getvalue()
